"""Levantamento de rede: o que um sensor enxerga, exportado em JSON (para IA e integrações) e XLSX.

O JSON é pensado para ser entregue a uma IA: cada informação inferida diz de onde veio e com que
confiança, e o próprio arquivo explica os campos ("sobre", "glossario", "limitacoes").
"""
import io
import unicodedata
from collections import Counter
from datetime import datetime, timezone

import oui

FORMATO = 'nexopulse.levantamento-rede'
VERSAO_FORMATO = 1

FONTES_NOME = {'dns': 'DNS reverso', 'netbios': 'NetBIOS', 'mdns': 'mDNS'}
FONTES_MARCA = {'ieee': 'Registro IEEE (prefixo do MAC)', 'sensor': 'Informado pelo sensor', 'mac_privado': 'MAC privado'}
CONFIANCA_MARCA = {'ieee': 'alta', 'sensor': 'media', 'mac_privado': 'nenhuma'}
STATUS = {'online': 'Responde a ping', 'sem_icmp': 'Presente na rede, não responde a ping',
          'desconhecido': 'Visto recentemente, sem confirmação', 'offline': 'Fora do ar'}

SOBRE = ('Levantamento dos equipamentos que o sensor NexoPulse enxerga na rede local dele (camada 2), '
         'gerado pela plataforma NexoPulse da MD Soluções. Cada dispositivo traz IP, MAC, a marca inferida '
         'pelo prefixo do MAC (registro público do IEEE) e o nome anunciado na rede, sempre com a fonte. '
         'Campos com valor null não foram identificados; não os trate como ausência do equipamento.')
LIMITACOES = [
    'Só aparecem equipamentos no mesmo segmento de rede (mesma VLAN/sub-rede) do sensor; '
    'equipamentos atrás de outros roteadores não entram nesta versão.',
    'A marca vem do prefixo do MAC: indica quem fabricou a placa de rede, que pode ser um fornecedor '
    'de módulos (ex.: Foxconn, AzureWave, Espressif) e não a marca do produto final.',
    'Celulares e notebooks atuais usam MAC privado (aleatório) por rede; nesses casos a marca não pode '
    'ser inferida pelo MAC.',
    'O nome vem do que o próprio equipamento ou o roteador anunciam (DNS reverso, NetBIOS, mDNS) e pode '
    'ter sido definido pelo usuário do equipamento.',
    'Equipamentos desligados ou que não falaram na rede durante a varredura não aparecem.',
]
GLOSSARIO = {
    'status': STATUS,
    'mac_privado': 'MAC com o bit "administrado localmente" ligado: gerado pelo sistema, não de fábrica.',
    'papel': {'gateway': 'Gateway padrão da rede do sensor', 'sensor': 'O próprio sensor NexoPulse'},
    'estado_vizinho': 'Estado da entrada na tabela de vizinhos do sensor (REACHABLE, STALE, DELAY...).',
}


def _agora():
    return datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')


def _chave_ip(ip):
    try:
        return (0, tuple(int(p) for p in str(ip).split('.')))
    except ValueError:
        return (1, str(ip))


def montar(sensor, dispositivos, versao_central='', ocultar_nomes=False):
    """Dicionário do levantamento a partir da linha do sensor e das linhas de dispositivos."""
    sensor = dict(sensor)
    ip_gw, ip_sensor = sensor.get('ip_gateway'), sensor.get('ip_sensor')
    itens = []
    vistos = set()
    for linha in sorted((dict(d) for d in dispositivos), key=lambda d: _chave_ip(d.get('ip'))):
        fonte_marca = linha.get('fabricante_fonte')
        apelido = linha.get('nome_custom')
        if apelido in ('Desconhecido', 'Gateway / Roteador'):
            apelido = None
        hostname = linha.get('hostname')
        papel = 'gateway' if linha.get('ip') == ip_gw else ('sensor' if linha.get('ip') == ip_sensor else None)
        vistos.add(linha.get('ip'))
        itens.append({
            'ip': linha.get('ip'),
            'mac': linha.get('mac'),
            'apelido': None if ocultar_nomes else apelido,
            'nome': {'valor': None if ocultar_nomes else hostname,
                     'fonte': FONTES_NOME.get(linha.get('hostname_fonte')) if hostname and not ocultar_nomes else None},
            'marca': {'valor': linha.get('fabricante') if fonte_marca in ('ieee', 'sensor') else None,
                      'razao_social': linha.get('fabricante_registro'),
                      'fonte': FONTES_MARCA.get(fonte_marca),
                      'confianca': CONFIANCA_MARCA.get(fonte_marca)},
            'mac_privado': fonte_marca == 'mac_privado',
            'papel': papel,
            'status': linha.get('status'),
            'latencia_ms': linha.get('latencia') if (linha.get('latencia') or 0) > 0 else None,
            'estado_vizinho': linha.get('neighbor_state') or None,
            'visto_em': linha.get('visto_em'),
        })
    if ip_sensor and ip_sensor not in vistos:
        # O sensor não aparece na própria tabela ARP; entra para o inventário ficar completo.
        mac = oui.normalizar(sensor.get('interface_mac'))
        registro, fonte = oui.fabricante(mac) if mac else (None, None)
        itens.append({
            'ip': ip_sensor, 'mac': mac,
            'apelido': None if ocultar_nomes else sensor.get('nome_local'),
            'nome': {'valor': None, 'fonte': None},
            'marca': {'valor': oui.marca(registro) if fonte == 'ieee' else None,
                      'razao_social': registro if fonte == 'ieee' else None,
                      'fonte': FONTES_MARCA.get(fonte), 'confianca': CONFIANCA_MARCA.get(fonte)},
            'mac_privado': fonte == 'mac_privado', 'papel': 'sensor', 'status': 'online',
            'latencia_ms': None, 'estado_vizinho': None, 'visto_em': None,
        })
        itens.sort(key=lambda d: _chave_ip(d['ip']))

    marcas = Counter(d['marca']['valor'] or ('MAC privado' if d['mac_privado'] else 'Não identificada') for d in itens)
    return {
        'formato': FORMATO,
        'versao_formato': VERSAO_FORMATO,
        'gerado_em': _agora(),
        'gerado_por': f'NexoPulse Central {versao_central}'.strip(),
        'sobre': SOBRE,
        'nomes_ocultos': bool(ocultar_nomes),
        'sensor': {
            'nome': sensor.get('nome_local'),
            'id': sensor.get('mac_id'),
            'cliente': sensor.get('cliente_nome'),
            'ip': ip_sensor,
            'gateway': ip_gw,
            'interface': sensor.get('interface_nome'),
            'provedor': sensor.get('provedor'),
            'versao_agente': sensor.get('agent_version'),
            'ultimo_contato': str(sensor['last_seen']) if sensor.get('last_seen') else None,
        },
        'escopo': {
            'rede': sensor.get('rede_cidr'),
            'faixa_varrida': sensor.get('scan_rede') or sensor.get('rede_cidr'),
            'varredura_limitada': bool(sensor.get('scan_limitado')),
            'metodo': ['Tabela de vizinhos (ARP) do sensor após varredura por ping',
                       'Marca pelo prefixo do MAC (registro IEEE MA-L, MA-M e MA-S)',
                       'Nome por DNS reverso, NetBIOS e mDNS'],
            'limitacoes': LIMITACOES,
        },
        'resumo': {
            'total': len(itens),
            'por_status': dict(Counter(d['status'] or 'desconhecido' for d in itens)),
            'por_marca': dict(marcas.most_common()),
            'com_nome': sum(1 for d in itens if d['nome']['valor'] or d['apelido']),
            'mac_privado': sum(1 for d in itens if d['mac_privado']),
        },
        'dispositivos': itens,
        'glossario': GLOSSARIO,
    }


# ---------------------------------------------------------------- XLSX
COLUNAS = (('IP', 16), ('MAC', 19), ('Apelido', 24), ('Nome na rede', 28), ('Fonte do nome', 14),
           ('Marca', 20), ('Fabricante (registro IEEE)', 38), ('MAC privado', 12), ('Papel', 10),
           ('Status', 34), ('Latência (ms)', 13), ('Estado L2', 12), ('Visto em (UTC)', 21))


def _linha_planilha(d):
    return [d['ip'], d['mac'], d['apelido'], d['nome']['valor'], d['nome']['fonte'], d['marca']['valor'],
            d['marca']['razao_social'], 'Sim' if d['mac_privado'] else 'Não',
            {'gateway': 'Gateway', 'sensor': 'Sensor'}.get(d['papel'], ''), STATUS.get(d['status'], d['status']),
            d['latencia_ms'], d['estado_vizinho'], d['visto_em']]


def _texto(celula, valor):
    """Grava texto sempre como texto: um nome de equipamento "=..." não pode virar fórmula."""
    celula.value = valor
    if isinstance(valor, str):
        celula.data_type = 's'


def para_xlsx(dados):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    cabecalho = Font(bold=True, color='FFFFFF')
    fundo = PatternFill('solid', fgColor='0B1B32')
    livro = Workbook()

    aba = livro.active
    aba.title = 'Dispositivos'
    for col, (titulo, largura) in enumerate(COLUNAS, start=1):
        celula = aba.cell(row=1, column=col)
        _texto(celula, titulo)
        celula.font, celula.fill = cabecalho, fundo
        aba.column_dimensions[get_column_letter(col)].width = largura
    for lin, d in enumerate(dados['dispositivos'], start=2):
        for col, valor in enumerate(_linha_planilha(d), start=1):
            _texto(aba.cell(row=lin, column=col), valor)
    aba.freeze_panes = 'A2'
    aba.auto_filter.ref = f'A1:{get_column_letter(len(COLUNAS))}{max(1, len(dados["dispositivos"]) + 1)}'

    resumo = livro.create_sheet('Resumo')
    s, e, r = dados['sensor'], dados['escopo'], dados['resumo']
    linhas = [('Levantamento de rede', None), ('Gerado em (UTC)', dados['gerado_em']), ('Sensor', s['nome']),
              ('Cliente', s['cliente']), ('IP do sensor', s['ip']), ('Gateway', s['gateway']), ('Rede', e['rede']),
              ('Faixa varrida', e['faixa_varrida']), ('Provedor', s['provedor']), ('Total de dispositivos', r['total']),
              ('Com nome identificado', r['com_nome']), ('Com MAC privado', r['mac_privado']), (None, None),
              ('Status', 'Quantidade')]
    linhas += [(STATUS.get(k, k), v) for k, v in r['por_status'].items()]
    linhas += [(None, None), ('Marca', 'Quantidade')] + list(r['por_marca'].items())
    for lin, (rotulo, valor) in enumerate(linhas, start=1):
        _texto(resumo.cell(row=lin, column=1), rotulo)
        _texto(resumo.cell(row=lin, column=2), valor)
        if rotulo in ('Levantamento de rede', 'Status', 'Marca'):
            for col in (1, 2):
                resumo.cell(row=lin, column=col).font = cabecalho
                resumo.cell(row=lin, column=col).fill = fundo
    resumo.column_dimensions['A'].width, resumo.column_dimensions['B'].width = 32, 40

    sobre = livro.create_sheet('Sobre')
    textos = [dados['sobre'], ''] + ['Limitações:'] + [f'• {t}' for t in e['limitacoes']]
    for lin, texto in enumerate(textos, start=1):
        celula = sobre.cell(row=lin, column=1)
        _texto(celula, texto)
        celula.alignment = Alignment(wrap_text=True, vertical='top')
    sobre.column_dimensions['A'].width = 110

    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def nome_arquivo(dados, extensao):
    # Só ASCII no cabeçalho HTTP: acentos viram a letra sem acento, o resto vira hífen.
    nome = unicodedata.normalize('NFKD', dados['sensor']['nome'] or 'sensor').encode('ascii', 'ignore').decode()
    base = ''.join(c if c.isalnum() or c in '-_' else '-' for c in nome)
    return f'levantamento-{base.strip("-")[:40] or "sensor"}-{dados["gerado_em"][:10]}.{extensao}'
